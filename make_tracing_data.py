"""Adds sheets 'Themes' (20) and 'TraceWords' (80) to data/worksheet_database.xlsx and extends the glossary.
Run once:  python make_tracing_data.py   (safe to re-run; it replaces those two sheets)."""
import csv
import re
import unicodedata
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

def slug(t):
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")

# Template, No, Theme_ID, Theme_EN, Tip_ID, Tip_EN, [(Word_ID, Word_EN, Syllables_ID override or '')]
T = [
("A","01","Peternakan","Farm","Minta anak menirukan suara dan gerakan hewannya, lalu ucapkan namanya.","Ask your child to imitate the animal's sound and movement, then say its name.",
 [("sapi","Cow",""),("ayam","Chicken",""),("domba","Sheep",""),("kuda","Horse","")]),
("B","02","Hewan Laut","Sea Animals","Tirukan gerakan berenang dengan tangan, lalu ucapkan nama hewannya.","Show a swimming movement with your hand, then say the animal's name.",
 [("hiu","Shark","hi-u"),("paus","Whale","pa-us"),("gurita","Octopus",""),("penyu","Turtle","")]),
("C","03","Buah-buahan","Fruits","Sebutkan warna dan rasa buahnya sambil menunjuk gambar.","Name the fruit's color and taste while pointing to the picture.",
 [("apel","Apple",""),("pisang","Banana",""),("jeruk","Orange",""),("anggur","Grapes","")]),
("D","04","Sayuran","Vegetables","Ajak anak menyebutkan sayuran yang pernah ia makan.","Ask your child which of these vegetables they have eaten.",
 [("wortel","Carrot",""),("tomat","Tomato",""),("bayam","Spinach",""),("jagung","Corn","")]),
("E","05","Kendaraan","Vehicles","Tirukan suara kendaraan, misalnya brum-brum, lalu ucapkan namanya.","Imitate the vehicle sound, for example vroom-vroom, then say its name.",
 [("mobil","Car",""),("bus","Bus",""),("sepeda","Bicycle",""),("kapal","Ship","")]),
("F","06","Anggota Tubuh","Body Parts","Sentuh bagian tubuh sendiri saat menyebut namanya.","Touch your own body part while saying its name.",
 [("mata","Eye",""),("hidung","Nose",""),("mulut","Mouth",""),("telinga","Ear","")]),
("G","07","Keluarga & Rumah","Family & Home","Tunjuk anggota keluarga di rumah, lalu sebutkan namanya.","Point to family members at home and say their names.",
 [("ayah","Father",""),("ibu","Mother",""),("nenek","Grandma",""),("rumah","House","")]),
("H","08","Perlengkapan Sekolah","School Things","Pegang benda aslinya bila ada, lalu ucapkan namanya.","Hold the real object if you have it, then say its name.",
 [("buku","Book",""),("pensil","Pencil",""),("tas","Bag",""),("penggaris","Ruler","")]),
("I","09","Cuaca & Alam","Weather & Nature","Lihat ke luar jendela dan tunjuk cuaca atau alam hari ini.","Look out of the window and point to today's weather or nature.",
 [("matahari","Sun",""),("hujan","Rain",""),("awan","Cloud",""),("pohon","Tree","")]),
("J","10","Makanan & Minuman","Food & Drink","Ajak anak menyebutkan makanan dan minuman kesukaannya.","Ask your child to name their favourite food and drink.",
 [("nasi","Rice",""),("roti","Bread",""),("susu","Milk",""),("air","Water","")]),
("K","11","Pakaian & Aksesori","Clothes","Minta anak menunjuk pakaian yang sedang ia pakai.","Ask your child to point to the clothes they are wearing.",
 [("baju","Shirt",""),("celana","Pants",""),("topi","Hat",""),("sepatu","Shoes","")]),
("L","12","Benda di Dapur","Kitchenware","Sebutkan benda dapur sambil membantu menyiapkan meja makan.","Name kitchen items while helping to set the table.",
 [("piring","Plate",""),("sendok","Spoon",""),("garpu","Fork",""),("gelas","Glass","")]),
("M","13","Profesi & Cita-Cita","Professions","Tanyakan: kamu ingin menjadi apa? Lalu tirukan pekerjaannya.","Ask: what do you want to be? Then act out the job.",
 [("dokter","Doctor",""),("guru","Teacher",""),("polisi","Police officer",""),("koki","Chef","")]),
("N","14","Benda di Rumah","Home Objects","Berkeliling rumah dan tunjuk benda yang sama.","Walk around the house and point to the same objects.",
 [("kursi","Chair",""),("meja","Table",""),("kasur","Bed",""),("lampu","Lamp","")]),
("O","15","Dinosaurus","Dinosaurs","Menderumlah seperti dinosaurus yang besar, lalu ucapkan namanya.","Stomp and roar like a big dinosaur, then say its name.",
 [("T-Rex","T-Rex",""),("Raptor","Raptor",""),("telur","Egg",""),("ekor","Tail","")]),
("P","16","Taman & Bunga","Garden & Flowers","Cium bunga asli bila ada dan sebutkan nama bagiannya.","Smell a real flower if you have one and name its parts.",
 [("bunga","Flower",""),("mawar","Rose",""),("kupu-kupu","Butterfly",""),("daun","Leaf","")]),
("Q","17","Dinosaurus & Purba","Prehistoric","Bicarakan bersama: benda dan hewan ini berasal dari zaman dahulu.","Talk about it together: these things are from a very long time ago.",
 [("fosil","Fossil",""),("mamut","Mammoth",""),("gua","Cave","gu-a"),("tulang","Bone","")]),
("R","18","Emosi & Perasaan","Feelings","Buat wajah sesuai perasaannya, lalu ucapkan namanya.","Make a face to match each feeling, then say its name.",
 [("senang","Happy",""),("sedih","Sad",""),("marah","Angry",""),("mengantuk","Sleepy","")]),
("S","19","Bentuk & Warna","Shapes & Colors","Cari benda berbentuk atau berwarna sama di sekitar kalian.","Find objects with the same shape or color around you.",
 [("lingkaran","Circle",""),("persegi","Square",""),("merah","Red",""),("biru","Blue","")]),
("T","20","Rasa & Sensori","Tastes & Sensations","Biarkan anak merasakan dengan aman: manis, asam, panas, dingin.","Let your child safely feel and name: sweet, sour, hot, cold.",
 [("manis","Sweet",""),("asam","Sour",""),("panas","Hot",""),("dingin","Cold","")]),
]
STYLE = "thick dark-navy outline, flat bright colors, plain white background, friendly smile, no text"

wb = load_workbook("data/worksheet_database.xlsx")
for n in ("Themes", "TraceWords"):
    if n in wb.sheetnames: del wb[n]
th = wb.create_sheet("Themes", 0)
tw = wb.create_sheet("TraceWords", 1)
th.append(["Template", "No", "Theme_ID", "Theme_EN", "Tip_ID", "Tip_EN"])
tw.append(["Template", "Word_ID", "Word_EN", "Syllables_ID", "Image_File", "Image_Prompt"])
for tpl, no, tid, ten, tipi, tipe, words in T:
    th.append([tpl, no, tid, ten, tipi, tipe])
    for wid, wen, syl in words:
        tw.append([tpl, wid, wen, syl, "", f"cute cartoon {wen.lower()} ({ten.lower()}), {STYLE}"])
for ws, widths in ((th, [10, 6, 24, 22, 70, 70]), (tw, [10, 18, 18, 18, 16, 90])):
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4D9DE0")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial"); c.number_format = "@"
    for col, w in zip("ABCDEF", widths): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
lg = wb["Legend"]
lg.append(["Sheet 'Themes'", "Tracing Words series (templates A-T): one row per page. Template = letter A-T. No = 01-20. Tip_ID / Tip_EN = parent tip printed at the bottom."])
lg.append(["Sheet 'TraceWords'", "Exactly 4 rows per Template. Image_File (optional) = file name in assets/images/ if you do not want the default {word_en}.png (e.g. cow.png, t_rex.png)."])
wb.save("data/worksheet_database.xlsx")

# ---- glossary: add EN->ID pairs
rows = list(csv.DictReader(open("data/vocabulary_glossary.csv", encoding="utf-8")))
idx = {r["EN"].lower(): r for r in rows}
for *_, words in T:
    for wid, wen, syl in words:
        k = wen.lower()
        if k in idx:
            if wid.lower() not in {x.strip().lower() for x in idx[k]["ID"].split("|")}:
                idx[k]["ID"] += "|" + wid.lower()
        else:
            r = {"EN": k, "ID": wid.lower(), "Countable": "Y", "Syllables_ID": syl}
            rows.append(r); idx[k] = r
with open("data/vocabulary_glossary.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=["EN", "ID", "Countable", "Syllables_ID"]); w.writeheader(); w.writerows(rows)
print("themes", len(T), "words", sum(len(t[-1]) for t in T), "glossary rows", len(rows))
