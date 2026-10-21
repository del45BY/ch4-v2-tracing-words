"""Creates data/worksheet_database.xlsx with 5 sample rows + a legend sheet."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = ["ID", "Category", "Object_EN", "Object_ID", "Relation_EN", "Relation_ID", "Image_Prompt"]
STYLE = "thick dark-navy outline, flat bright colors, plain white background, friendly smile, no text"
ROWS = [
    ["001", "Animals & Their Products", "Chicken", "Ayam", "Egg", "Telur", f"cute cartoon chicken holding a big egg, {STYLE}"],
    ["002", "Animals & Their Products", "Bee", "Lebah", "Honey", "Madu", f"cute cartoon bee next to a honey jar, {STYLE}"],
    ["003", "Animals & Their Products", "Cow", "Sapi", "Milk", "Susu", f"cute cartoon cow face beside a milk carton, {STYLE}"],
    ["004", "Animals & Their Products", "Sheep", "Domba", "Wool", "Wol", f"cute cartoon fluffy sheep beside a ball of yarn, {STYLE}"],
    ["005", "Plants & Fruits", "Tree", "Pohon", "Apple", "Apel", f"cute cartoon apple tree and one big red apple, {STYLE}"],
]
wb = Workbook()
ws = wb.active
ws.title = "Worksheets"
ws.append(HEAD)
for r in ROWS:
    ws.append(r)
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4D9DE0")
    c.alignment = Alignment(vertical="center")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial")
        c.number_format = "@"
for col, w in zip("ABCDEFG", [7, 28, 14, 14, 14, 14, 90]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
dv = DataValidation(type="textLength", operator="between", formula1="1", formula2="30", showErrorMessage=True,
                    errorTitle="Too long", error="Keep words short (max 30 characters).")
ws.add_data_validation(dv)
dv.add("C2:F1000")

lg = wb.create_sheet("Legend")
lines = [
    ("Column", "Meaning"),
    ("ID", "Unique 3-digit code (001, 002 ...). Used for ordering and for repeatable randomisation."),
    ("Category", "Folder name AND wording group. Add new groups in wsengine/config.py CATEGORY_WORDING."),
    ("Object_EN / Object_ID", "Main picture word. English then Indonesian. Singular, simple, 3-10 letters is ideal."),
    ("Relation_EN / Relation_ID", "The paired picture word (what the object gives / needs / is next to)."),
    ("Image_Prompt", "Prompt for AI art. Artwork is saved as assets/images/{object_en}.png (lowercase, e.g. chicken.png)."),
    ("Tip", "Every word must exist in data/vocabulary_glossary.csv or QA will warn you for a human check."),
]
for l in lines:
    lg.append(list(l))
for c in lg[1]:
    c.font = Font(name="Arial", bold=True)
lg.column_dimensions["A"].width = 28
lg.column_dimensions["B"].width = 110
for row in lg.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial")
wb.save("data/worksheet_database.xlsx")
print("saved data/worksheet_database.xlsx")
